"""Excel exports with live formulas.

valuation_workbook(result)  -> every valuation model on its own sheet, all inputs on an Inputs sheet, formulas wired so
                                changing an assumption re-computes the sheet, sensitivity grid and blend.
technicals_workbook(result, columns) -> contiguous Bars table (date, OHLCV, chosen indicators) so "select all → Insert chart"
                                works, plus pre-built native Excel charts for the price/overlays and each oscillator.
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F3A5F")
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
SUB = Font(italic=True, color="666666")
BOLD = Font(bold=True)


def _hdr(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill = HDR, FILL
        cell.alignment = Alignment(horizontal="center")
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _g(d: Any, *path, default=None):
    cur = d
    for p in path:
        if not isinstance(cur, dict):
            return default
        cur = cur.get(p)
        if cur is None:
            return default
    return cur


def _num(x, default=0.0):
    try:
        v = float(x)
        return default if v != v else v
    except (TypeError, ValueError):
        return default


# =============================================================================== valuation
def valuation_workbook(res: dict[str, Any]) -> bytes:
    v = res.get("valuation") or {}
    inp = v.get("inputs") or {}
    wb = Workbook()

    # ---------------- Inputs (named cells)
    ws = wb.active
    ws.title = "Inputs"
    _hdr(ws, 1, ["Assumption", "Value", "Source", "Name"], [34, 16, 60, 18])
    rows = [
        ("Symbol", res.get("symbol"), "", "Symbol"), ("Price", _g(inp, "price", "value"), _g(inp, "price", "source"), "Price"),
        ("Shares outstanding", _g(inp, "shares", "value"), _g(inp, "shares", "source"), "Shares"), ("Net debt", _g(inp, "net_debt", "value"), _g(inp, "net_debt", "source"), "NetDebt"),
        ("Revenue (TTM)", _g(inp, "revenue_ttm", "value"), _g(inp, "revenue_ttm", "source"), "Revenue"), ("Risk-free rate", _g(inp, "rf", "value"), _g(inp, "rf", "source"), "Rf"),
        ("Equity risk premium", _g(inp, "erp", "value"), _g(inp, "erp", "source"), "ERP"), ("Beta", _g(inp, "beta", "value"), _g(inp, "beta", "source"), "Beta"),
        ("Cost of equity", "=Rf+Beta*ERP", "rf + beta × ERP", "Ke"), ("Pre-tax cost of debt", _g(inp, "cost_of_debt_pre_tax", "value"), _g(inp, "cost_of_debt_pre_tax", "source"), "Kd"),
        ("Tax rate", _g(inp, "tax_rate", "value"), _g(inp, "tax_rate", "source"), "TaxRate"), ("Weight equity", _g(inp, "weight_equity", "value"), _g(inp, "weight_equity", "source"), "We"),
        ("Weight debt", "=1-We", "", "Wd"), ("WACC", "=We*Ke+Wd*Kd*(1-TaxRate)", "we·ke + wd·kd·(1−t)", "WACC"),
        ("Stage-1 revenue growth", _g(inp, "growth_stage1", "value"), _g(inp, "growth_stage1", "source"), "GrowthS1"), ("Terminal growth", _g(inp, "growth_terminal", "value"), _g(inp, "growth_terminal", "source"), "GrowthT"),
        ("Fade years", _g(inp, "fade_years", "value", default=10), "growth fades linearly to terminal", "FadeYrs"),
        ("FCF margin (today)", _g(v, "dcf", "margin", default=_g(inp, "fcf_margin", "value")), _g(inp, "fcf_margin", "source"), "MarginNow"),
        ("FCF margin (target, year 10)", _g(v, "dcf", "margin_target", default=_g(inp, "fcf_margin", "value")), "normalized target when today's margin is negative", "MarginTarget"),
        ("EBITDA margin", _g(inp, "extra", "ebitda_margin"), "TTM", "EbitdaM"), ("EBIT margin", _g(inp, "ebit_margin", "value"), _g(inp, "ebit_margin", "source"), "EbitM"),
        ("Exit EV/EBITDA multiple", _g(v, "dcf_exit", "exit_multiple", default=12), _g(v, "dcf_exit", "exit_multiple_source"), "ExitMult"),
        ("EPS (TTM)", _g(inp, "eps_ttm", "value"), _g(inp, "eps_ttm", "source"), "EPS"), ("Dividend / share", _g(inp, "dps", "value"), _g(inp, "dps", "source"), "DPS"),
        ("Book value / share", _g(inp, "bvps", "value"), _g(inp, "bvps", "source"), "BVPS"), ("ROE", _g(inp, "roe", "value"), _g(inp, "roe", "source"), "ROE"),
        ("Payout ratio", _g(inp, "payout", "value"), _g(inp, "payout", "source"), "Payout"),
    ]
    for i, (label, val, src, name) in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=label)
        c = ws.cell(row=i, column=2, value=val)
        c.fill = INPUT_FILL
        ws.cell(row=i, column=3, value=src).font = SUB
        ws.cell(row=i, column=4, value=name).font = SUB
        wb.defined_names[name] = DefinedName(name, attr_text=f"Inputs!$B${i}")
        if isinstance(val, float) and abs(val) < 5 and label not in ("Price", "Beta", "EPS (TTM)", "Dividend / share", "Book value / share", "Fade years", "Exit EV/EBITDA multiple"):
            c.number_format = "0.00%"
        elif isinstance(val, (int, float)):
            c.number_format = "#,##0.00"
    ws.cell(row=len(rows) + 3, column=1, value="Yellow cells are inputs — change them and every sheet recomputes. Prices are delayed; profile: " + str(_g(v, "profile", "primary"))).font = SUB

    # ---------------- DCF
    d = wb.create_sheet("DCF")
    _hdr(d, 1, ["Year", "Growth", "Margin", "Revenue", "FCF", "Discount factor", "PV(FCF)"], [8, 12, 12, 18, 18, 16, 18])
    d.cell(row=2, column=1, value=0); d.cell(row=2, column=4, value="=Revenue"); d.cell(row=2, column=3, value="=MarginNow")
    for y in range(1, 11):
        r = y + 2
        d.cell(row=r, column=1, value=y)
        d.cell(row=r, column=2, value=f"=GrowthS1+(GrowthT-GrowthS1)*MIN(A{r}-1,FadeYrs)/FadeYrs")
        d.cell(row=r, column=3, value=f"=MarginNow+(MarginTarget-MarginNow)*MIN(A{r},FadeYrs)/FadeYrs")
        d.cell(row=r, column=4, value=f"=D{r - 1}*(1+B{r})")
        d.cell(row=r, column=5, value=f"=D{r}*C{r}")
        d.cell(row=r, column=6, value=f"=1/(1+WACC)^A{r}")
        d.cell(row=r, column=7, value=f"=E{r}*F{r}")
        for col in (2, 3):
            d.cell(row=r, column=col).number_format = "0.00%"
        for col in (4, 5, 7):
            d.cell(row=r, column=col).number_format = "#,##0"
    out = [("PV of FCF (years 1-10)", "=SUM(G3:G12)"), ("Terminal value (Gordon)", "=E12*(1+GrowthT)/(WACC-GrowthT)"), ("PV of terminal value", "=B15*F12"),
           ("Enterprise value", "=B14+B16"), ("Terminal share of EV", "=B16/B17"), ("Equity value", "=B17-NetDebt"), ("Value per share (Gordon)", "=B19/Shares"),
           ("Upside vs price", "=B20/Price-1"), ("", ""), ("Terminal value (exit multiple)", "=D12*EbitdaM*ExitMult"), ("PV of exit terminal", "=B23*F12"),
           ("Value per share (exit multiple)", "=(B14+B24-NetDebt)/Shares"), ("Upside vs price (exit)", "=B25/Price-1")]
    for i, (label, f) in enumerate(out, 14):
        d.cell(row=i, column=1, value=label).font = BOLD
        c = d.cell(row=i, column=2, value=f)
        c.number_format = "0.00%" if "Upside" in label or "share of" in label else "#,##0.00"
    d.column_dimensions["A"].width = 30
    wb.defined_names["DCF_PerShare"] = DefinedName("DCF_PerShare", attr_text="DCF!$B$20")
    wb.defined_names["DCFExit_PerShare"] = DefinedName("DCFExit_PerShare", attr_text="DCF!$B$25")

    # ---------------- Sensitivity (WACC × terminal g) on the same projection
    s = wb.create_sheet("Sensitivity")
    s.cell(row=1, column=1, value="Value per share: WACC (rows) × terminal growth (columns); FCF path from the DCF sheet").font = BOLD
    dw = [-0.02, -0.01, 0, 0.01, 0.02]
    dg = [-0.01, -0.005, 0, 0.005, 0.01]
    for j, g in enumerate(dg):
        c = s.cell(row=3, column=2 + j, value=f"=GrowthT+({g})")
        c.number_format, c.font, c.fill = "0.00%", HDR, FILL
    for i, w in enumerate(dw):
        r = 4 + i
        c = s.cell(row=r, column=1, value=f"=WACC+({w})")
        c.number_format, c.font, c.fill = "0.00%", HDR, FILL
        for j in range(len(dg)):
            col = get_column_letter(2 + j)
            f = f"=(SUMPRODUCT(DCF!$E$3:$E$12,1/(1+$A{r})^DCF!$A$3:$A$12)+DCF!$E$12*(1+{col}$3)/($A{r}-{col}$3)/(1+$A{r})^10-NetDebt)/Shares"
            s.cell(row=r, column=2 + j, value=f).number_format = "#,##0.00"
    s.column_dimensions["A"].width = 12

    # ---------------- Scenarios
    sc = wb.create_sheet("Scenarios")
    _hdr(sc, 1, ["", "Bear", "Base", "Bull"], [28, 16, 16, 16])
    cases = _g(v, "scenarios", "cases", default={}) or {}
    params = [("Probability", [_g(cases, k, "probability", default=p) for k, p in (("bear", 0.25), ("base", 0.5), ("bull", 0.25))]),
              ("Stage-1 growth", ["=GrowthS1-0.08", "=GrowthS1", "=GrowthS1+0.08"]), ("Target margin", ["=MarginTarget*0.8", "=MarginTarget", "=MarginTarget*1.2"]), ("WACC", ["=WACC+0.01", "=WACC", "=MAX(WACC-0.01,GrowthT+0.01)"])]
    for i, (label, vals) in enumerate(params, 2):
        sc.cell(row=i, column=1, value=label).font = BOLD
        for j, val in enumerate(vals):
            c = sc.cell(row=i, column=2 + j, value=val)
            c.fill = INPUT_FILL
            c.number_format = "0.00%"
    # explicit revenue rows (8-17) and FCF rows (20-29) per scenario — growth fades to GrowthT, margin ramps from MarginNow to the scenario target
    sc.cell(row=7, column=1, value="Revenue by year").font = BOLD
    sc.cell(row=19, column=1, value="FCF by year").font = BOLD
    for y in range(1, 11):
        r, rf_ = 7 + y, 19 + y
        sc.cell(row=r, column=1, value=y); sc.cell(row=rf_, column=1, value=y)
        for j, col in enumerate(("B", "C", "D")):
            prev = "Revenue" if y == 1 else f"{col}{r - 1}"
            sc.cell(row=r, column=2 + j, value=f"={prev}*(1+{col}$3+(GrowthT-{col}$3)*MIN($A{r}-1,FadeYrs)/FadeYrs)").number_format = "#,##0"
            sc.cell(row=rf_, column=2 + j, value=f"={col}{r}*(MarginNow+({col}$4-MarginNow)*MIN($A{rf_},FadeYrs)/FadeYrs)").number_format = "#,##0"
    sc.cell(row=31, column=1, value="Value per share").font = BOLD
    for j, col in enumerate(("B", "C", "D")):
        sc.cell(row=31, column=2 + j, value=f"=(SUMPRODUCT({col}20:{col}29,1/(1+{col}$5)^$A$20:$A$29)+{col}29*(1+GrowthT)/({col}$5-GrowthT)/(1+{col}$5)^10-NetDebt)/Shares").number_format = "#,##0.00"
    sc.cell(row=32, column=1, value="Probability-weighted value").font = BOLD
    sc.cell(row=32, column=2, value="=SUMPRODUCT(B2:D2,B31:D31)").number_format = "#,##0.00"
    wb.defined_names["Scenario_PerShare"] = DefinedName("Scenario_PerShare", attr_text="Scenarios!$B$32")

    # ---------------- Earnings / book / dividend models
    m = wb.create_sheet("Models")
    _hdr(m, 1, ["Model", "Value / share", "Formula"], [34, 16, 70])
    models = [
        ("Earnings power value (EPV)", "=IF(EbitM>0,(Revenue*EbitM*(1-TaxRate)/WACC-NetDebt)/Shares,\"n/a\")", "EBIT×(1−t)/WACC − net debt, per share"),
        ("Dividend discount (Gordon)", "=IF(DPS>0,DPS*(1+MIN(MAX(ROE*(1-Payout),0),Ke-0.01,0.08))/(Ke-MIN(MAX(ROE*(1-Payout),0),Ke-0.01,0.08)),\"n/a\")", "DPS×(1+g)/(ke−g), g = min(ROE×(1−payout), ke−1%, 8%)"),
        ("Justified P/E × EPS", "=IF(AND(EPS>0,Payout>0),Payout*(1+MIN(MAX(GrowthT,0),Ke-0.01))/(Ke-MIN(MAX(GrowthT,0),Ke-0.01))*EPS,\"n/a\")", "payout×(1+g)/(ke−g) × EPS"),
        ("Justified P/B × BVPS", "=IF(AND(BVPS>0,ROE>MIN(MAX(GrowthT,0),Ke-0.01)),(ROE-MIN(MAX(GrowthT,0),Ke-0.01))/(Ke-MIN(MAX(GrowthT,0),Ke-0.01))*BVPS,\"n/a\")", "(ROE−g)/(ke−g) × BVPS"),
        ("Graham number", "=IF(AND(EPS>0,BVPS>0),SQRT(22.5*EPS*BVPS),\"n/a\")", "√(22.5 × EPS × BVPS)"),
        ("PEG-based (PEG = 1)", "=IF(AND(EPS>0,GrowthS1>0.03),MIN(GrowthS1*100,50)*EPS,\"n/a\")", "P/E = growth% × 1, capped 50×"),
        ("Residual income (5y fade)", None, "computed in the Residual Income sheet"),
    ]
    for i, (label, f, desc) in enumerate(models, 2):
        m.cell(row=i, column=1, value=label).font = BOLD
        c = m.cell(row=i, column=2, value=f if f else "=ResidualIncome!B12")
        c.number_format = "#,##0.00"
        m.cell(row=i, column=3, value=desc).font = SUB
    # residual income sheet
    ri = wb.create_sheet("ResidualIncome")
    _hdr(ri, 1, ["Year", "Opening BV", "ROE (fading)", "Residual income", "PV"], [8, 16, 14, 16, 16])
    for y in range(1, 6):
        r = y + 1
        ri.cell(row=r, column=1, value=y)
        ri.cell(row=r, column=2, value="=BVPS" if y == 1 else f"=B{r - 1}*(1+C{r - 1}*(1-Payout))")
        ri.cell(row=r, column=3, value=f"=ROE-(ROE-Ke)*({y - 1})/5").number_format = "0.00%"
        ri.cell(row=r, column=4, value=f"=(C{r}-Ke)*B{r}")
        ri.cell(row=r, column=5, value=f"=D{r}/(1+Ke)^A{r}")
    ri.cell(row=8, column=1, value="Terminal (half-strength RI, years 6-10)").font = BOLD
    ri.cell(row=8, column=5, value="=(ROE-Ke)*B6*(1+C6*(1-Payout))*0.5*SUMPRODUCT(1/(1+Ke)^{6,7,8,9,10})")
    ri.cell(row=12, column=1, value="Value per share").font = BOLD
    ri.cell(row=12, column=2, value="=IF(BVPS>0,BVPS+SUM(E2:E6)+E8,\"n/a\")").number_format = "#,##0.00"

    # ---------------- Comps
    cp = wb.create_sheet("Comps")
    comps = v.get("comps") or {}
    peers = comps.get("peers") or []
    cols = ["symbol", "name", "market_cap", "rev_growth", "gross_margin", "pe", "fwd_pe", "ev_ebitda", "ev_sales", "pb", "peg"]
    _hdr(cp, 1, cols, [10, 28, 16, 12, 12, 10, 10, 12, 10, 10, 10])
    for i, p in enumerate(peers, 2):
        for j, c in enumerate(cols, 1):
            cp.cell(row=i, column=j, value=p.get(c))
    n = len(peers)
    if n:
        rr = n + 3
        cp.cell(row=rr, column=1, value="median").font = BOLD
        cp.cell(row=rr + 1, column=1, value="subject").font = BOLD
        cp.cell(row=rr + 2, column=1, value="implied price @ median").font = BOLD
        stats = comps.get("stats") or {}
        for j, c in enumerate(cols, 1):
            if c in ("pe", "fwd_pe", "ev_ebitda", "ev_sales", "pb", "peg"):
                L = get_column_letter(j)
                cp.cell(row=rr, column=j, value=f"=MEDIAN({L}2:{L}{n + 1})").number_format = "0.0"
                subj = _g(stats, c, "subject")
                cp.cell(row=rr + 1, column=j, value=subj).fill = INPUT_FILL
                cp.cell(row=rr + 2, column=j, value=f"=IF(AND(ISNUMBER({L}{rr + 1}),{L}{rr + 1}>0),Price*{L}{rr}/{L}{rr + 1},\"\")").number_format = "#,##0.00"
        cp.cell(row=rr + 4, column=1, value="Blended implied (median of implied prices)").font = BOLD
        cp.cell(row=rr + 4, column=2, value=f"=MEDIAN(F{rr + 2}:K{rr + 2})").number_format = "#,##0.00"
        wb.defined_names["Comps_PerShare"] = DefinedName("Comps_PerShare", attr_text=f"Comps!$B${rr + 4}")
    else:
        cp.cell(row=2, column=1, value="no peer data in this run (quick depth)")

    # ---------------- Historical multiples
    hm = wb.create_sheet("HistMultiples")
    hist = v.get("hist_multiples_history") or []
    _hdr(hm, 1, ["fy_end", "price", "pe", "ps", "ev_ebitda", "pb"], [14, 12, 10, 10, 12, 10])
    for i, h in enumerate(hist, 2):
        for j, c in enumerate(["fy_end", "price", "pe", "ps", "ev_ebitda", "pb"], 1):
            hm.cell(row=i, column=j, value=h.get(c))
    if hist:
        rr = len(hist) + 3
        hm.cell(row=rr, column=1, value="median").font = BOLD
        for j, c in enumerate(["pe", "ps", "ev_ebitda", "pb"], 3):
            L = get_column_letter(j)
            hm.cell(row=rr, column=j, value=f"=MEDIAN({L}2:{L}{len(hist) + 1})").number_format = "0.0"
        ttm = {"eps": _g(inp, "eps_ttm", "value"), "rev": _g(inp, "revenue_ttm", "value"), "ebitda": _g(inp, "extra", "ebitda_ttm"), "equity": _g(inp, "extra", "equity_ttm")}
        hm.cell(row=rr + 1, column=1, value="TTM basis").font = BOLD
        hm.cell(row=rr + 1, column=3, value=ttm["eps"]); hm.cell(row=rr + 1, column=4, value=ttm["rev"]); hm.cell(row=rr + 1, column=5, value=ttm["ebitda"]); hm.cell(row=rr + 1, column=6, value=ttm["equity"])
        hm.cell(row=rr + 2, column=1, value="implied / share").font = BOLD
        hm.cell(row=rr + 2, column=3, value=f"=IF(C{rr + 1}>0,C{rr}*C{rr + 1},\"\")"); hm.cell(row=rr + 2, column=4, value=f"=IF(D{rr + 1}>0,D{rr}*D{rr + 1}/Shares,\"\")")
        hm.cell(row=rr + 2, column=5, value=f"=IF(E{rr + 1}>0,(E{rr}*E{rr + 1}-NetDebt)/Shares,\"\")"); hm.cell(row=rr + 2, column=6, value=f"=IF(F{rr + 1}>0,F{rr}*F{rr + 1}/Shares,\"\")")
        hm.cell(row=rr + 3, column=1, value="median implied").font = BOLD
        hm.cell(row=rr + 3, column=2, value=f"=MEDIAN(C{rr + 2}:F{rr + 2})").number_format = "#,##0.00"
        wb.defined_names["Hist_PerShare"] = DefinedName("Hist_PerShare", attr_text=f"HistMultiples!$B${rr + 3}")

    # ---------------- Blend
    b = wb.create_sheet("Blend")
    _hdr(b, 1, ["Model", "Value / share", "Weight", "Applies?", "Weighted"], [34, 16, 10, 10, 16])
    blend = v.get("blend") or {}
    refs = {"dcf": "=DCF_PerShare", "dcf_exit": "=DCFExit_PerShare", "dcf_scenarios": "=Scenario_PerShare", "comps": "=IFERROR(Comps_PerShare,\"\")", "hist_multiples": "=IFERROR(Hist_PerShare,\"\")",
            "epv": "=Models!B2", "ddm": "=Models!B3", "justified_pe": "=Models!B4", "justified_pb": "=Models!B5", "graham": "=Models!B6", "peg": "=Models!B7", "rim": "=Models!B8"}
    row = 2
    for mm in blend.get("models") or []:
        key = mm["key"]
        b.cell(row=row, column=1, value=mm["label"])
        val = refs.get(key, mm.get("value_per_share"))
        b.cell(row=row, column=2, value=val).number_format = "#,##0.00"
        b.cell(row=row, column=3, value=mm.get("weight") or 0).number_format = "0%"
        b.cell(row=row, column=3).fill = INPUT_FILL
        b.cell(row=row, column=4, value="yes" if mm.get("applies") else "no")
        b.cell(row=row, column=5, value=f"=IF(AND(ISNUMBER(B{row}),D{row}=\"yes\"),B{row}*C{row},0)").number_format = "#,##0.00"
        row += 1
    b.cell(row=row + 1, column=1, value="Blended fair value").font = BOLD
    b.cell(row=row + 1, column=2, value=f"=SUM(E2:E{row - 1})/SUMPRODUCT((D2:D{row - 1}=\"yes\")*ISNUMBER(B2:B{row - 1})*C2:C{row - 1})").number_format = "#,##0.00"
    b.cell(row=row + 2, column=1, value="Upside vs price").font = BOLD
    b.cell(row=row + 2, column=2, value=f"=B{row + 1}/Price-1").number_format = "0.00%"
    b.cell(row=row + 4, column=1, value=f"Profile: {blend.get('profile')} · weights are the profile defaults; edit the yellow cells to re-weight.").font = SUB

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =============================================================================== technicals
def technicals_workbook(res: dict[str, Any], overlay_cols: list[str], pane_cols: list[list[str]]) -> bytes:
    t = res.get("technicals") or {}
    frame = t.get("indicator_frame_1d") or []
    bars = t.get("bars", {}).get("1d") or []
    by_ts = {r["ts"]: r for r in bars}
    wb = Workbook()
    ws = wb.active
    ws.title = "Bars"
    cols = ["ts", "open", "high", "low", "close", "volume"] + [c for c in overlay_cols if frame and c in frame[0]] + [c for grp in pane_cols for c in grp if frame and c in frame[0]]
    seen: list[str] = []
    for c in cols:
        if c not in seen:
            seen.append(c)
    cols = seen
    _hdr(ws, 1, cols, [12] + [12] * (len(cols) - 1))
    for i, r in enumerate(frame, 2):
        b = by_ts.get(r["ts"], {})
        for j, c in enumerate(cols, 1):
            val = b.get(c, r.get(c)) if c in ("open", "high", "low", "volume") else r.get(c)
            if c == "ts":
                val = str(val)[:10]
            ws.cell(row=i, column=j, value=val)
    n = len(frame) + 1
    ws.freeze_panes = "B2"
    # native charts: price + overlays, then one per pane
    def line_chart(title, col_names, anchor, y_title):
        ch = LineChart()
        ch.title, ch.y_axis.title, ch.x_axis.title = title, y_title, "date"
        ch.height, ch.width = 9, 24
        for cn in col_names:
            if cn in cols:
                j = cols.index(cn) + 1
                ch.add_data(Reference(ws, min_col=j, min_row=1, max_row=n), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
        ws2.add_chart(ch, anchor)
    ws2 = wb.create_sheet("Charts")
    line_chart("Close + overlays", ["close"] + [c for c in overlay_cols if c in cols], "A1", "price")
    row = 20
    for grp in pane_cols:
        gs = [c for c in grp if c in cols]
        if gs:
            line_chart(" / ".join(gs), gs, f"A{row}", "value")
            row += 19
    # latest values sheet
    lv = wb.create_sheet("Latest")
    _hdr(lv, 1, ["indicator", "value"], [28, 16])
    latest = t.get("intervals", {}).get("1d", {}).get("latest") or {}
    for i, (k, v) in enumerate(sorted(latest.items()), 2):
        lv.cell(row=i, column=1, value=k)
        lv.cell(row=i, column=2, value=v)
    sig = t.get("intervals", {}).get("1d", {}).get("signals") or {}
    r0 = len(latest) + 4
    lv.cell(row=r0, column=1, value="signals").font = BOLD
    for i, (k, v) in enumerate(sig.items(), r0 + 1):
        lv.cell(row=i, column=1, value=k)
        lv.cell(row=i, column=2, value=str(v))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
